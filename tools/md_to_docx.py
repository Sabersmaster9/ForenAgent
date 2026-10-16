"""把交付包里的 Markdown 转成 Word（.docx），供老师和专利代理人审阅。

支持：标题、段落、表格、代码块、有序/无序列表、引用、粗体与行内代码。

用法：
    D:/CondaData/envs/forenagent/python.exe tools/md_to_docx.py <目录或文件> [...]
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CJK_BODY = "宋体"
CJK_HEAD = "黑体"
MONO = "Consolas"

HEADING_LEVELS = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6}


def _set_cjk(run, cjk: str = CJK_BODY, mono: bool = False) -> None:
    run.font.name = MONO if mono else cjk
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), MONO if mono else cjk)


_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(paragraph, text: str) -> None:
    """写入行内格式：**粗体** 与 `代码`。"""
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
            _set_cjk(run)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.font.size = Pt(9.5)
            _set_cjk(run, mono=True)
        else:
            run = paragraph.add_run(piece)
            _set_cjk(run)


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and len(s) > 2


def _split_row(line: str):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = CJK_BODY
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_BODY)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(12)
            para.paragraph_format.space_after = Pt(6)
            run = para.add_run("\n".join(block))
            run.font.size = Pt(9)
            _set_cjk(run, mono=True)
            continue

        # 表格
        if _is_table_row(line):
            rows = []
            while i < len(lines) and _is_table_row(lines[i]):
                if not _is_separator(lines[i]):
                    rows.append(_split_row(lines[i]))
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    cells = table.add_row().cells
                    for c_idx in range(width):
                        text = row[c_idx] if c_idx < len(row) else ""
                        para = cells[c_idx].paragraphs[0]
                        para.text = ""
                        add_runs(para, text)
                        for run in para.runs:
                            run.font.size = Pt(9)
                            if r_idx == 0:
                                run.bold = True
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = HEADING_LEVELS[len(m.group(1))]
            heading = doc.add_heading("", level=level)
            add_runs(heading, m.group(2))
            for run in heading.runs:
                _set_cjk(run, CJK_HEAD)
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run("— — —")
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            _set_cjk(run)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(18)
            add_runs(para, stripped.lstrip("> ").strip())
            for run in para.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # 列表
        m_ul = re.match(r"^[-*]\s+(.*)$", stripped)
        m_ol = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        m_cb = re.match(r"^-\s*\[([ xX])\]\s*(.*)$", stripped)
        if m_cb:
            para = doc.add_paragraph(style="List Bullet")
            add_runs(para, ("☑ " if m_cb.group(1).lower() == "x" else "☐ ") + m_cb.group(2))
            i += 1
            continue
        if m_ul:
            para = doc.add_paragraph(style="List Bullet")
            add_runs(para, m_ul.group(1))
            i += 1
            continue
        if m_ol:
            para = doc.add_paragraph(style="List Number")
            add_runs(para, m_ol.group(2))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        para = doc.add_paragraph()
        add_runs(para, stripped)
        para.paragraph_format.space_after = Pt(6)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main(argv: list) -> int:
    targets = []
    for arg in argv[1:]:
        p = Path(arg)
        if p.is_dir():
            targets.extend(sorted(p.rglob("*.md")))
        elif p.suffix.lower() == ".md":
            targets.append(p)
    if not targets:
        print("未找到 Markdown 文件")
        return 1
    for md in targets:
        out = md.with_suffix(".docx")
        convert(md, out)
        print(f"  {out.name}  ({out.stat().st_size // 1024} KB)")
    print(f"共生成 {len(targets)} 个 .docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
